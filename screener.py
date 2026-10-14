"""Core logic: extract resume text and score it against a job description with Gemini."""

import io
import json
import os
from typing import List, Literal

from docx import Document
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader

DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

SYSTEM_PROMPT = """You are an objective recruiting assistant that screens resumes against a job description.

Rules:
- Judge ONLY on skills, experience, education, and achievements relevant to the role.
- Do NOT consider or infer name, gender, age, ethnicity, nationality, religion, marital status,
  photo, or any other protected characteristic. Do not penalize career gaps without evidence they matter.
- Base every claim on evidence in the resume. If something is not stated, treat it as unknown, not as missing proof of weakness.
- Be calibrated: 90+ = exceptional fit, 70-89 = strong, 50-69 = partial, below 50 = weak.
- Your output is decision support for a human recruiter, not a final hiring decision."""


class Screening(BaseModel):
    """Structured screening result for one resume."""

    candidate_name: str = Field(description="Name as written on the resume, or 'Unknown'.")
    score: int = Field(ge=0, le=100, description="Overall fit score.")
    recommendation: Literal["Strong Match", "Potential Match", "Weak Match"]
    summary: str = Field(description="2-3 sentence rationale.")
    years_experience: float = Field(description="Approx. years of relevant experience; 0 if unclear.")
    matched_skills: List[str]
    missing_skills: List[str]
    strengths: List[str]
    concerns: List[str]
    interview_questions: List[str] = Field(description="3-5 tailored questions.")


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX, or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()
    raise ValueError(f"Unsupported file type: {filename}")


def get_client() -> genai.Client:
    key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not key:
        raise RuntimeError("GEMINI_API_KEY is not set. Add it to your .env file.")
    return genai.Client(api_key=key)


def screen_resume(client: genai.Client, job_description: str, resume_text: str,
                  model: str = DEFAULT_MODEL, must_haves: str = "") -> dict:
    """Score one resume against the job description. Returns the structured result dict."""
    if len(resume_text) < 50:
        raise ValueError("Resume text is empty or too short (scanned image PDF? Try OCR first).")

    contents = (
        f"<job_description>\n{job_description}\n</job_description>\n\n"
        + (f"<must_have_requirements>\n{must_haves}\n</must_have_requirements>\n\n" if must_haves.strip() else "")
        + f"<resume>\n{resume_text[:30000]}\n</resume>\n\n"
        "Screen this resume against the job description."
    )

    response = client.models.generate_content(
        model=model,
        contents=contents,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            response_mime_type="application/json",
            response_schema=Screening,
            temperature=0.2,
        ),
    )

    if isinstance(response.parsed, Screening):
        return response.parsed.model_dump()
    if response.text:
        return Screening(**json.loads(response.text)).model_dump()
    raise RuntimeError("Gemini returned no result (the response may have been blocked).")
